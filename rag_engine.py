"""
RAG Engine: PDF 추출 → 청킹 → 벡터DB → Gemini 챗봇
"""
import os, re, json, hashlib, time
from typing import List, Dict, Optional, Tuple
import google.generativeai as genai
import chromadb
from chromadb.utils import embedding_functions

# ── Gemini 모델 설정 ──
FLASH_MODELS = [
    "gemini-2.5-flash-preview-05-20",
    "gemini-2.0-flash-001",
    "gemini-1.5-flash-latest",
]
EMBED_MODEL = "models/text-embedding-004"


def configure_gemini(api_key: str):
    """Gemini API 키 설정"""
    genai.configure(api_key=api_key)


def get_flash_model() -> genai.GenerativeModel:
    """사용 가능한 Flash 모델 반환 (Cascading fallback)"""
    for model_name in FLASH_MODELS:
        try:
            model = genai.GenerativeModel(model_name)
            model.generate_content("test", generation_config={"max_output_tokens": 5})
            return model
        except Exception:
            continue
    raise RuntimeError("사용 가능한 Gemini Flash 모델이 없습니다.")


# ============================================================
# 1. PDF 텍스트 추출
# ============================================================
def extract_text_from_pdf(file_bytes: bytes) -> str:
    """PDF 바이트에서 텍스트 추출 (PyMuPDF)"""
    import fitz
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    texts = []
    for page_num, page in enumerate(doc):
        text = page.get_text("text")
        if text.strip():
            texts.append(f"[페이지 {page_num + 1}]\n{text.strip()}")
    doc.close()
    return "\n\n".join(texts)


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """파일 형식에 따라 텍스트 추출"""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")
    elif ext == ".docx":
        from docx import Document
        import io
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    else:
        return file_bytes.decode("utf-8", errors="ignore")


# ============================================================
# 2. 텍스트 청킹
# ============================================================
def chunk_text(text: str, chunk_size: int = 800, overlap: int = 150) -> List[Dict]:
    """텍스트를 오버랩 청킹"""
    paragraphs = re.split(r'\n{2,}', text)
    chunks = []
    current_chunk = ""
    current_page = "1"

    for para in paragraphs:
        # 페이지 번호 추출
        page_match = re.match(r'\[페이지 (\d+)\]', para)
        if page_match:
            current_page = page_match.group(1)
            para = re.sub(r'\[페이지 \d+\]\n?', '', para)

        if len(current_chunk) + len(para) > chunk_size and current_chunk:
            chunks.append({
                "text": current_chunk.strip(),
                "page": current_page,
            })
            # 오버랩: 마지막 일부 유지
            words = current_chunk.split()
            overlap_words = words[-overlap // 4:] if len(words) > overlap // 4 else words
            current_chunk = " ".join(overlap_words) + "\n" + para
        else:
            current_chunk += "\n" + para if current_chunk else para

    if current_chunk.strip():
        chunks.append({
            "text": current_chunk.strip(),
            "page": current_page,
        })
    return chunks


# ============================================================
# 3. Gemini 토픽 태깅
# ============================================================
def tag_topics_batch(chunks: List[Dict], category: str) -> List[Dict]:
    """청크 배치에 토픽 태그 자동 부여 (Gemini)"""
    topic_maps = {
        "macro": ["경제지표", "인구사회", "건강웰니스", "기술디지털", "규제정책", "환경ESG", "음료산업", "소비트렌드", "기타"],
        "beverage": ["탄산음료", "비탄산음료", "기능성음료", "RTD커피차", "유제품음료", "식물성음료", "주류", "시장규모", "브랜드분석", "유통채널", "기타"],
        "consumer": ["구매동기", "연령별", "성별", "채널별", "가격민감도", "브랜드충성도", "건강관심", "트렌드인식", "기타"],
    }
    valid_topics = topic_maps.get(category, ["기타"])

    model = get_flash_model()
    tagged_chunks = []

    # 배치 처리 (10개씩)
    for i in range(0, len(chunks), 10):
        batch = chunks[i:i + 10]
        batch_texts = "\n---\n".join(
            [f"[CHUNK_{j}] {c['text'][:300]}" for j, c in enumerate(batch)]
        )

        prompt = f"""다음 텍스트 청크들에 토픽 태그를 부여하세요.
가능한 토픽: {', '.join(valid_topics)}

각 청크에 1~3개의 가장 관련 높은 토픽을 선택하세요.
JSON 배열로만 응답하세요. 다른 텍스트 없이.

형식: [{{"chunk": 0, "topics": ["토픽1", "토픽2"]}}, ...]

텍스트:
{batch_texts}"""

        try:
            resp = model.generate_content(prompt, generation_config={
                "temperature": 0.1,
                "max_output_tokens": 1024,
            })
            text = resp.text.strip()
            # JSON 파싱
            text = re.sub(r'^```json\s*', '', text)
            text = re.sub(r'\s*```$', '', text)
            results = json.loads(text)
            topic_map = {r["chunk"]: r["topics"] for r in results}
        except Exception:
            topic_map = {}

        for j, chunk in enumerate(batch):
            topics = topic_map.get(j, ["기타"])
            chunk["topics"] = topics
            chunk["category"] = category
            tagged_chunks.append(chunk)

        time.sleep(0.3)  # Rate limit

    return tagged_chunks


# ============================================================
# 4. ChromaDB 벡터 저장소
# ============================================================
class RAGVectorStore:
    """ChromaDB 기반 벡터 저장소"""

    def __init__(self, api_key: str):
        self.client = chromadb.Client()
        self.embed_fn = embedding_functions.GoogleGenerativeAiEmbeddingFunction(
            api_key=api_key,
            model_name=EMBED_MODEL,
        )
        self.collections = {}

    def get_or_create_collection(self, name: str) -> chromadb.Collection:
        if name not in self.collections:
            self.collections[name] = self.client.get_or_create_collection(
                name=name,
                embedding_function=self.embed_fn,
                metadata={"hnsw:space": "cosine"},
            )
        return self.collections[name]

    def add_chunks(self, collection_name: str, chunks: List[Dict]):
        """청크들을 컬렉션에 추가"""
        coll = self.get_or_create_collection(collection_name)

        ids = []
        documents = []
        metadatas = []

        for i, chunk in enumerate(chunks):
            chunk_id = hashlib.md5(
                f"{collection_name}_{i}_{chunk['text'][:50]}".encode()
            ).hexdigest()
            ids.append(chunk_id)
            documents.append(chunk["text"])
            metadatas.append({
                "category": chunk.get("category", ""),
                "topics": ",".join(chunk.get("topics", ["기타"])),
                "page": chunk.get("page", ""),
            })

        # 배치 추가 (100개씩)
        for start in range(0, len(ids), 100):
            end = min(start + 100, len(ids))
            coll.add(
                ids=ids[start:end],
                documents=documents[start:end],
                metadatas=metadatas[start:end],
            )

    def query(
        self,
        collection_name: str,
        query_text: str,
        n_results: int = 5,
        topic_filter: Optional[str] = None,
    ) -> List[Dict]:
        """벡터 검색 + 토픽 필터"""
        coll = self.get_or_create_collection(collection_name)
        if coll.count() == 0:
            return []

        where_filter = None
        if topic_filter and topic_filter != "전체":
            where_filter = {"topics": {"$contains": topic_filter}}

        try:
            results = coll.query(
                query_texts=[query_text],
                n_results=min(n_results, coll.count()),
                where=where_filter,
            )
        except Exception:
            # 필터 실패 시 필터 없이 재시도
            results = coll.query(
                query_texts=[query_text],
                n_results=min(n_results, coll.count()),
            )

        docs = []
        if results and results["documents"]:
            for doc, meta, dist in zip(
                results["documents"][0],
                results["metadatas"][0],
                results["distances"][0],
            ):
                docs.append({
                    "text": doc,
                    "metadata": meta,
                    "distance": dist,
                })
        return docs

    def query_multiple(
        self,
        collection_names: List[str],
        query_text: str,
        n_results: int = 5,
    ) -> List[Dict]:
        """여러 컬렉션 동시 검색"""
        all_docs = []
        for name in collection_names:
            docs = self.query(name, query_text, n_results=n_results)
            all_docs.extend(docs)
        # 거리 기준 정렬 후 상위 반환
        all_docs.sort(key=lambda x: x["distance"])
        return all_docs[:n_results * 2]

    def clear_collection(self, name: str):
        """컬렉션 초기화"""
        try:
            self.client.delete_collection(name)
            if name in self.collections:
                del self.collections[name]
        except Exception:
            pass


# ============================================================
# 5. Gemini RAG 챗봇
# ============================================================
SYSTEM_PROMPTS = {
    "macro": """당신은 거시시장 분석 전문가입니다. 
제공된 보고서 내용을 기반으로 정확하게 답변하세요.
데이터와 수치를 인용하고, 출처 페이지를 표시하세요.
음료·식품 산업 관점의 시사점을 우선적으로 도출하세요.
답변 마지막에 "📌 보고서에 '[섹션제목]' 섹션 추가됨"을 포함하세요.""",

    "beverage": """당신은 음료시장 분석 전문가입니다.
제공된 보고서 내용을 기반으로 정확하게 답변하세요.
시장규모, 성장률, 브랜드, 유통채널 등 구체적 데이터를 포함하세요.
경쟁 환경과 기회 영역을 식별하세요.
답변 마지막에 "📌 보고서에 '[섹션제목]' 섹션 추가됨"을 포함하세요.""",

    "consumer": """당신은 소비자 리서치 분석 전문가입니다.
제공된 조사 데이터를 기반으로 정확하게 답변하세요.
수치(%, N수)를 반드시 인용하고, 연령/성별/채널별 세분화 분석을 하세요.
마케팅 시사점을 함께 도출하세요.
답변 마지막에 "📌 보고서에 '[섹션제목]' 섹션 추가됨"을 포함하세요.""",

    "formula": """당신은 음료 배합 설계 전문가입니다.
제공된 원료 DB, 가이드 배합비, 트렌드 분석 결과를 기반으로 배합비를 설계하세요.
반드시 포함할 정보: 원료명, 배합비(%), 예상 Brix, pH, 산도, 감미도, 칼로리, 원가
DB의 원료 물성치(1%사용시 기여도)를 활용해 계산하세요.
답변 마지막에 "📌 배합비 업데이트됨"을 포함하세요.""",

    "report": """당신은 식품·음료 산업 전략 컨설턴트입니다.
거시시장, 음료시장, 소비자 태도 3개 분석 결과와 음료 DB를 종합하여 크로스 분석하세요.
인과관계를 명확히 하고(거시→시장→소비자), 전략적 기회 영역을 도출하세요.
트렌드에서 구체적 제품 컨셉과 배합비까지 연결하세요.
답변 마지막에 "📌 보고서에 '[섹션제목]' 섹션 추가됨"을 포함하세요.""",
}


def build_rag_prompt(
    query: str,
    retrieved_docs: List[Dict],
    system_prompt: str,
    chat_history: List[Dict] = None,
) -> str:
    """RAG 프롬프트 구성"""
    context_parts = []
    for i, doc in enumerate(retrieved_docs):
        meta = doc.get("metadata", {})
        topics = meta.get("topics", "")
        page = meta.get("page", "")
        context_parts.append(
            f"[참조 {i+1}] (토픽: {topics}, 페이지: {page})\n{doc['text']}"
        )

    context_text = "\n\n---\n\n".join(context_parts) if context_parts else "(검색된 관련 문서 없음)"

    # 대화 히스토리
    history_text = ""
    if chat_history:
        recent = chat_history[-6:]  # 최근 3턴
        for msg in recent:
            role = "사용자" if msg["role"] == "user" else "AI"
            history_text += f"\n{role}: {msg['content']}"

    prompt = f"""{system_prompt}

═══ 참조 문서 ═══
{context_text}

═══ 대화 기록 ═══
{history_text}

═══ 사용자 질문 ═══
{query}

위 참조 문서를 기반으로 정확하고 구조화된 답변을 제공하세요.
보고서에 추가할 수 있는 형태로 작성하세요."""

    return prompt


def chat_with_rag(
    query: str,
    vector_store: "RAGVectorStore",
    collection_name: str,
    chat_history: List[Dict],
    topic_filter: str = "전체",
    extra_collections: List[str] = None,
) -> str:
    """RAG 기반 Gemini 챗봇 응답 생성"""
    # 1. 검색
    if extra_collections:
        all_names = [collection_name] + extra_collections
        retrieved = vector_store.query_multiple(all_names, query, n_results=6)
    else:
        retrieved = vector_store.query(
            collection_name, query, n_results=6, topic_filter=topic_filter
        )

    # 2. 프롬프트 구성
    system_prompt = SYSTEM_PROMPTS.get(collection_name, SYSTEM_PROMPTS["macro"])
    prompt = build_rag_prompt(query, retrieved, system_prompt, chat_history)

    # 3. Gemini 생성
    model = get_flash_model()
    response = model.generate_content(
        prompt,
        generation_config={
            "temperature": 0.4,
            "max_output_tokens": 2048,
        },
    )
    return response.text


# ============================================================
# 6. 자동 요약 생성
# ============================================================
def generate_auto_summary(text: str, category: str) -> Tuple[str, List[str]]:
    """PDF 전체 텍스트 자동 요약 + 키워드 추출"""
    model = get_flash_model()

    # 텍스트 길이 제한 (토큰 절약)
    max_chars = 30000
    if len(text) > max_chars:
        text = text[:max_chars // 2] + "\n...(중략)...\n" + text[-max_chars // 2:]

    category_names = {"macro": "거시시장", "beverage": "음료시장", "consumer": "소비자 태도"}
    cat_name = category_names.get(category, category)

    prompt = f"""다음 {cat_name} 보고서를 분석하여:

1. 핵심 요약 (3~5문장, 주요 수치 포함)
2. 핵심 키워드 6~8개

JSON 형식으로만 응답하세요:
{{"summary": "요약 텍스트", "keywords": ["키워드1", "키워드2", ...]}}

보고서 내용:
{text}"""

    try:
        resp = model.generate_content(prompt, generation_config={
            "temperature": 0.2,
            "max_output_tokens": 512,
        })
        result_text = resp.text.strip()
        result_text = re.sub(r'^```json\s*', '', result_text)
        result_text = re.sub(r'\s*```$', '', result_text)
        data = json.loads(result_text)
        return data.get("summary", "요약 생성 실패"), data.get("keywords", [])
    except Exception as e:
        return f"요약 생성 중 오류: {str(e)}", []


# ============================================================
# 7. 보고서 섹션 자동 추출
# ============================================================
def extract_report_section(ai_response: str, section_number: str) -> Optional[Dict]:
    """AI 응답에서 보고서 섹션을 추출"""
    model = get_flash_model()

    prompt = f"""다음 AI 분석 응답을 보고서 섹션으로 변환하세요.

AI 응답:
{ai_response}

JSON으로만 응답:
{{"title": "섹션 제목", "content": "보고서에 넣을 정리된 내용 (불릿포인트, 수치 포함)"}}"""

    try:
        resp = model.generate_content(prompt, generation_config={
            "temperature": 0.1,
            "max_output_tokens": 1024,
        })
        result_text = resp.text.strip()
        result_text = re.sub(r'^```json\s*', '', result_text)
        result_text = re.sub(r'\s*```$', '', result_text)
        data = json.loads(result_text)
        return {
            "number": section_number,
            "title": data.get("title", "제목 없음"),
            "content": data.get("content", ai_response),
            "is_new": True,
        }
    except Exception:
        # 폴백: 응답 그대로 사용
        title_match = re.search(r'📌 보고서에 [\'"](.+?)[\'"]', ai_response)
        title = title_match.group(1) if title_match else "분석 결과"
        return {
            "number": section_number,
            "title": title,
            "content": ai_response,
            "is_new": True,
        }


# ============================================================
# 8. 음료 DB 로드 (Excel)
# ============================================================
def load_beverage_db(file_bytes: bytes) -> Dict:
    """음료 DB 엑셀 파일 로드"""
    import pandas as pd
    import io

    xls = pd.ExcelFile(io.BytesIO(file_bytes))
    db = {}
    for sheet in xls.sheet_names:
        try:
            df = pd.read_excel(xls, sheet_name=sheet)
            db[sheet] = df
        except Exception:
            continue
    return db


def beverage_db_to_text(db: Dict) -> str:
    """음료 DB를 RAG용 텍스트로 변환"""
    texts = []

    # 원료DB
    if "원료DB" in db:
        df = db["원료DB"]
        for _, row in df.iterrows():
            parts = []
            for col in df.columns:
                val = row[col]
                if pd.notna(val):
                    parts.append(f"{col}: {val}")
            texts.append(" | ".join(parts))

    # 가이드배합비DB
    if "가이드배합비DB" in db:
        df = db["가이드배합비DB"]
        for _, row in df.iterrows():
            parts = []
            for col in df.columns:
                val = row[col]
                if pd.notna(val):
                    parts.append(f"{col}: {val}")
            texts.append(" | ".join(parts))

    # 시장제품DB
    if "시장제품DB" in db:
        df = db["시장제품DB"]
        for _, row in df.iterrows():
            parts = []
            for col in df.columns[:10]:  # 주요 컬럼만
                val = row[col]
                if pd.notna(val):
                    parts.append(f"{col}: {val}")
            texts.append(" | ".join(parts))

    return "\n".join(texts)


# ============================================================
# 9. PDF 보고서 생성 (fpdf2)
# ============================================================
def generate_pdf_report(
    title: str,
    sections: List[Dict],
    filename: str = "report.pdf",
) -> bytes:
    """보고서를 PDF로 생성"""
    from fpdf import FPDF
    import io

    class KoreanPDF(FPDF):
        def header(self):
            self.set_fill_color(37, 99, 235)
            self.rect(0, 0, 210, 15, 'F')
            self.set_font("Helvetica", "B", 9)
            self.set_text_color(255, 255, 255)
            self.set_y(4)
            self.cell(0, 7, "Trend Analyzer Report", align="C")
            self.ln(15)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = KoreanPDF()
    pdf.alias_nb_pages()

    # 한글 폰트 시도 (없으면 기본 폰트)
    font_paths = [
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansKR-Regular.ttf",
    ]
    korean_font = None
    for fp in font_paths:
        if os.path.exists(fp):
            try:
                pdf.add_font("Korean", "", fp, uni=True)
                pdf.add_font("Korean", "B", fp.replace("Regular", "Bold").replace("Gothic", "GothicBold"), uni=True)
                korean_font = "Korean"
                break
            except Exception:
                continue

    font = korean_font or "Helvetica"

    # 표지
    pdf.add_page()
    pdf.ln(40)
    pdf.set_font(font, "B", 24)
    pdf.set_text_color(30, 41, 59)
    pdf.multi_cell(0, 12, title, align="C")
    pdf.ln(10)
    pdf.set_font(font, "", 12)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 8, f"Gemini 2.5 Flash RAG Analysis | {time.strftime('%Y.%m.%d')}", align="C")

    # 섹션들
    for section in sections:
        pdf.add_page()
        # 섹션 번호 + 제목
        pdf.set_font(font, "B", 16)
        pdf.set_text_color(37, 99, 235)
        pdf.cell(0, 10, f"{section['number']}. {section['title']}")
        pdf.ln(12)
        # 내용
        pdf.set_font(font, "", 11)
        pdf.set_text_color(51, 65, 85)
        pdf.multi_cell(0, 6, section["content"])

    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


# ============================================================
# 10. 통합 분석 프롬프트
# ============================================================
def generate_integrated_analysis(
    macro_summary: str,
    beverage_summary: str,
    consumer_summary: str,
    query: str,
) -> str:
    """3개 보고서 크로스 분석"""
    model = get_flash_model()

    prompt = f"""{SYSTEM_PROMPTS['report']}

═══ 거시시장 분석 요약 ═══
{macro_summary}

═══ 음료시장 분석 요약 ═══
{beverage_summary}

═══ 소비자 태도 분석 요약 ═══
{consumer_summary}

═══ 질문 ═══
{query}

세 보고서를 크로스 분석하여 인과관계와 기회 영역을 도출하세요."""

    response = model.generate_content(
        prompt,
        generation_config={
            "temperature": 0.4,
            "max_output_tokens": 2048,
        },
    )
    return response.text


# pandas import for beverage_db_to_text
try:
    import pandas as pd
except ImportError:
    pass
